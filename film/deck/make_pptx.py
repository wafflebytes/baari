"""The venue .pptx: every slide's 4K final frame, full bleed, with the talk track in the
speaker notes (DECK_PLAN section 6). The videos go in where they play in the deck once
they land: the trailer before slide 1, the tour inside slide 8's phone, the night on
slide 9. Until then those slides carry their final frame, as in the PDF.

  node export.mjs --png --4k && node export.mjs --pdf   (makes out/4k/*.jpg)
  python3 make_pptx.py                                      (writes out/Baari_finale.pptx)
"""
import json
import os
import subprocess

from pptx import Presentation
from pptx.util import Emu

HERE = os.path.dirname(os.path.abspath(__file__))
EMU_PX = 6350  # 12,192,000 EMU (13.333 in, the 16:9 default) over 1920 px


def px(v):
    return Emu(round(v * EMU_PX))


# timing.js and talk.js are the source; node prints them as JSON
data = json.loads(subprocess.check_output(["node", "-e", """
global.window = {};
require('./timing.js'); require('./talk.js');
console.log(JSON.stringify({ T: window.DECK_TIMING, TALK: window.DECK_TALK }));
"""], cwd=HERE))
T, TALK = data["T"], data["TALK"]

# where each video plays, in slide pixels: s08's tour sits in the phone's screen below the
# status bar (the README frame at 400 px), s09's night plays full bleed
VIDEOS = {
    "trailer": ("../trailer/renders/trailer.mp4", None),
    "s08": ("videos/tour.mp4", (511, 219, 354, 767)),
    "s09": ("videos/night.mp4", (0, 0, 1920, 1080)),
}

prs = Presentation()
prs.slide_width, prs.slide_height = px(1920), px(1080)
blank = prs.slide_layouts[6]
missing = []


def page(img, notes):
    s = prs.slides.add_slide(blank)
    s.shapes.add_picture(img, 0, 0, prs.slide_width, prs.slide_height)
    s.notes_slide.notes_text_frame.text = notes
    return s


trailer = os.path.join(HERE, VIDEOS["trailer"][0])
if os.path.exists(trailer):
    s = prs.slides.add_slide(blank)
    s.shapes.add_movie(trailer, 0, 0, prs.slide_width, prs.slide_height, poster_frame_image=os.path.join(HERE, "out/4k/s01.jpg"), mime_type="video/mp4")
    s.notes_slide.notes_text_frame.text = "The trailer. Let it play, then go on."
else:
    missing.append("trailer")

for sl in T["slides"]:
    sid = sl["id"]
    img = os.path.join(HERE, f"out/4k/{sid}.jpg")
    if not os.path.exists(img):
        raise SystemExit(f"missing {img}: run node export.mjs --png --4k, then --pdf")
    talk = TALK.get(sid, {})
    who = f"{talk['who']}. " if talk.get("who") else ""
    timing = f"[{sl['talk']} s] " if sl.get("talk") else ""
    notes = f"{timing}{who}{talk.get('say', '')}"
    if sid in VIDEOS:
        path, (x, y, w, h) = os.path.join(HERE, VIDEOS[sid][0]), VIDEOS[sid][1]
        if os.path.exists(path):
            # the slide as it plays (its frame at 2 s, the video in its frame), then its final page
            bg = os.path.join(HERE, f"out/4k/{sid}-t2.png")
            if not os.path.exists(bg):
                subprocess.check_call(["node", "export.mjs", "--4k", "--at", sid, "2"], cwd=HERE)
            poster = os.path.join(HERE, f"out/4k/{sid}-poster.jpg")
            subprocess.check_call(["ffmpeg", "-y", "-v", "error", "-ss", "0.5", "-i", path, "-frames:v", "1", poster])
            s = page(bg, notes)
            s.shapes.add_movie(path, px(x), px(y), px(w), px(h), poster_frame_image=poster, mime_type="video/mp4")
            page(img, "The page this slide ends on, for the PDF reader. Go on.")
            continue
        missing.append(sid)
    page(img, notes)

out = os.path.join(HERE, "out/Baari_finale.pptx")
prs.save(out)
print(out, len(prs.slides), "slides", f"(no video yet for {', '.join(missing)}: final frame instead)" if missing else "")
