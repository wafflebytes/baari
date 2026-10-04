"""Build evals/out/baari_run_log.xlsx, the run log for answer 9 (EVAL_PLAN 5.2).

Import it into Google Sheets (File > Import > Upload, "Replace spreadsheet")
and every tab comes across. Sources: out/runs.csv, open_coding.md,
m1_models.md and agent/prompts/CHANGELOG.md. Rerun after each round.
"""
import csv
import os
import re
from collections import OrderedDict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

HERE = os.path.dirname(os.path.abspath(__file__))
EVALS = os.path.dirname(HERE)
ROOT = os.path.dirname(EVALS)


def md_tables(path):
    """Every markdown table in a file, each with the heading above it."""
    out, heading, rows = [], "", []
    for line in open(path, encoding="utf-8"):
        line = line.rstrip("\n")
        if line.startswith("#"):
            heading = line.lstrip("#").strip()
        if line.startswith("|"):
            cells = [c.strip() for c in line.strip().strip("|").split("|")]
            if all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
                continue
            rows.append(cells)
        elif rows:
            out.append((heading, rows))
            rows = []
    if rows:
        out.append((heading, rows))
    return out


def sheet(wb, title, header, rows, widths=None):
    ws = wb.create_sheet(title)
    ws.append(header)
    for r in rows:
        ws.append(r)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="0D0D0E")
    ws.freeze_panes = "A2"
    for i, w in enumerate(widths or []):
        ws.column_dimensions[chr(65 + i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    return ws


def tables_sheet(wb, title, path):
    ws = wb.create_sheet(title)
    for heading, rows in md_tables(path):
        if heading:
            ws.append([heading])
            ws.cell(ws.max_row, 1).font = Font(bold=True, size=12)
        ws.append(rows[0])
        for c in ws[ws.max_row]:
            c.font = Font(bold=True)
        for r in rows[1:]:
            ws.append(r)
        ws.append([])
    for col in "ABCDEFG":
        ws.column_dimensions[col].width = 38
    for row in ws.iter_rows():
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")


def main():
    runs = list(csv.DictReader(open(os.path.join(EVALS, "out", "runs.csv"), encoding="utf-8")))
    wb = Workbook()
    wb.remove(wb.active)

    rounds = OrderedDict()
    for r in runs:
        if r["target"] != "platform":
            continue
        k = (r["round"], r["prompt"], r["model"])
        g = rounds.setdefault(k, {"first": r["at"], "last": r["at"], "runs": 0, "pass": 0, "cases": OrderedDict()})
        g["last"] = r["at"]
        g["runs"] += 1
        g["pass"] += r["pass"] == "pass"
        g["cases"][r["case"]] = r["pass"]
    sheet(wb, "Rounds", ["Round", "Prompt", "Model", "First run (UTC)", "Last run (UTC)", "Runs", "Runs passed", "Cases", "Latest result per case passing"],
          [[k[0], k[1], k[2], g["first"], g["last"], g["runs"], g["pass"], " ".join(g["cases"]), f'{sum(v == "pass" for v in g["cases"].values())} of {len(g["cases"])}'] for k, g in rounds.items()],
          [8, 8, 34, 24, 24, 7, 11, 40, 16])

    cols = ["round", "case", "model", "prompt", "input", "at", "platform_run_id", "pass", "fails", "first_fail", "tool_calls", "ms", "trace"]
    sheet(wb, "Runs", [c.replace("_", " ").capitalize() for c in cols], [[r[c] for c in cols] for r in runs if r["target"] == "platform"],
          [7, 6, 32, 7, 10, 24, 38, 6, 6, 60, 8, 8, 50])
    replica = [r for r in runs if r["target"] != "platform"]
    if replica:
        sheet(wb, "Off-platform (replica)", [c.replace("_", " ").capitalize() for c in cols], [[r[c] for c in cols] for r in replica],
              [7, 6, 32, 7, 10, 24, 38, 6, 6, 60, 8, 8, 50])

    tables_sheet(wb, "Models", os.path.join(EVALS, "m1_models.md"))
    tables_sheet(wb, "Failures", os.path.join(EVALS, "open_coding.md"))

    pv = md_tables(os.path.join(ROOT, "agent", "prompts", "CHANGELOG.md"))
    rows = pv[0][1] if pv else [["Version"]]
    sheet(wb, "Prompt versions", rows[0], rows[1:], [14, 22, 8, 90, 30, 40])

    out = os.path.join(EVALS, "out", "baari_run_log.xlsx")
    wb.save(out)
    print(out, {ws.title: ws.max_row for ws in wb.worksheets})


if __name__ == "__main__":
    main()
